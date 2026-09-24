#!/usr/bin/env python3
"""
Generate the synthetic benchmark dataset: five fictional Indian private
companies with statements, notes, an auditor's report and an expense-voucher
ledger, plus planted issues, QA questions and ground truth.

Everything is derived from the specifications in this file with fixed seeds,
so a re-run produces the same content. No real company data is used; every
name, figure and event is invented.

Usage:
    python benchmarks/generate_dataset.py            # writes benchmarks/data/

Per company (benchmarks/data/<slug>/):
    annual_report.md      directors' report extract, statements, notes, auditor's report
    annual_report.pdf     the same content as a PDF
    financials.xlsx       balance sheet, P&L and related-party schedule (Rs. lakhs)
    general_ledger.csv    expense-voucher register for FY 2024-25 (rupees)

Shared:
    benchmarks/data/ground_truth.json   figures, planted issues and decoys per company
    benchmarks/data/qa.json             30 questions with gold answers
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import math
import random
import sys
import zipfile
from pathlib import Path

DATA_DIR = Path(__file__).resolve().parent / "data"
TAX_RATE = 0.2517  # section 115BAA effective rate
FY_START = dt.date(2024, 4, 1)
FY_END = dt.date(2025, 3, 31)
CUR, PRI = "FY 2024-25", "FY 2023-24"
BS_CUR, BS_PRI = "31 Mar 2025", "31 Mar 2024"
DISCLAIMER = (
    "Fictional company created for a benchmark. All names, figures and events "
    "in this document are invented."
)

# Issue taxonomy used by every system in the issue-detection benchmark
ISSUE_TYPES = {
    "rpt_threshold_no_approval": (
        "A related-party transaction above the Companies Act materiality threshold "
        "(e.g. 10% of net worth or turnover) without the required shareholder approval, "
        "or reported as not requiring approval"
    ),
    "rp_omitted_from_schedule": (
        "Transactions with a party related to a director that are missing from the "
        "related-party disclosure note"
    ),
    "duplicate_entries": "The same voucher or payment recorded more than once",
    "round_number_entries": "An unusually high share of round-number amounts in the ledger",
    "benford_nonconformity": (
        "Ledger amounts whose leading digits deviate from Benford's law, e.g. many "
        "amounts clustered just below an approval limit"
    ),
    "caro_missing": (
        "The auditor's report omits the Companies (Auditor's Report) Order (CARO) "
        "statement although the company is not exempt"
    ),
    "ratio_anomaly": (
        "A financial ratio that moves abnormally year on year (e.g. gross margin, "
        "receivable days) without a matching change in the business"
    ),
    "cash_loan_269ss": "A loan or deposit of Rs. 20,000 or more accepted in cash (section 269SS)",
    "unusual_large_payment": "A single payment far larger than the rest of the ledger",
}

# ---------------------------------------------------------------------------
# Company specifications (amounts in Rs. lakhs unless stated)
# ---------------------------------------------------------------------------
PNL_KEYS = [
    ("revenue", "Revenue from operations"),
    ("other_income", "Other income"),
    ("materials", "Cost of materials consumed"),
    ("employee", "Employee benefits expense"),
    ("finance", "Finance costs"),
    ("depreciation", "Depreciation and amortisation expense"),
    ("other_exp", "Other expenses"),
]
BS_KEYS = [
    ("share_capital", "Equity share capital"),
    ("lt_borrowings", "Long-term borrowings"),
    ("st_borrowings", "Short-term borrowings"),
    ("trade_payables", "Trade payables"),
    ("other_cl", "Other current liabilities"),
    ("ppe", "Property, plant and equipment"),
    ("inventories", "Inventories"),
    ("receivables", "Trade receivables"),
    ("cash", "Cash and cash equivalents"),
]

COMPANIES: list[dict] = [
    {
        "slug": "kestrel",
        "name": "Kestrel Polymers Private Limited",
        "short": "Kestrel Polymers",
        "city": "Nashik", "state": "Maharashtra", "incorporated": 2009,
        "business": "manufactures injection-moulded plastic components for appliance makers",
        "directors": [("Rohan Kulkarni", "Managing Director"), ("Sneha Kulkarni", "Director"),
                      ("Vivek Joshi", "Director")],
        "auditor": ("Sahyadri Audit Partners LLP", "Meera Gokhale", "Pune"),
        "report_date": "28 August 2025",
        "approval_limit": 100000,
        "pnl": {"revenue": (4386.40, 3978.25), "other_income": (42.15, 36.80),
                "materials": (2548.70, 2301.45), "employee": (612.35, 561.20),
                "finance": (118.60, 124.90), "depreciation": (205.45, 188.30),
                "other_exp": (698.45, 452.10)},
        "bs": {"share_capital": (400.00, 400.00), "other_equity_prior": 1512.40,
               "lt_borrowings": (640.00, 720.00), "st_borrowings": (385.50, 342.00),
               "trade_payables": (512.30, 468.75), "other_cl": (164.20, 151.60),
               "ppe": (1640.60, 1705.30), "inventories": (604.80, 548.20),
               "receivables": (812.45, 735.10), "cash": (148.95, 160.20)},
        "other_exp_lines": [("Power and fuel", 0.30), ("Rent", None), ("Repairs and maintenance", 0.10),
                            ("Freight and forwarding", 0.12), ("Legal and professional fees", None),
                            ("Travelling and conveyance", 0.05), ("Insurance", 0.03),
                            ("Security charges", None)],
        "fixed_lines": {"Rent": (9.60, 9.60), "Security charges": (7.20, 6.60),
                        "Legal and professional fees": (212.40, 24.65)},
        "footnotes": [
            "Legal and professional fees include Rs. 185.00 lakhs paid to Zenith Advisory Services "
            "for strategic consultancy.",
            "Rent is paid to Mrs. Sneha Kulkarni, a director, for the office premises at Nashik (see Note 27).",
        ],
        "mbp1": [("Rohan Kulkarni", "Kestrel Estates LLP", "Designated partner"),
                 ("Rohan Kulkarni", "Kulkarni Resins LLP", "Partner"),
                 ("Vivek Joshi", "Godavari Tooling Private Limited", "Director")],
        "rp_parties": [("Rohan Kulkarni", "Managing Director (key managerial personnel)"),
                       ("Sneha Kulkarni", "Director; wife of the Managing Director"),
                       ("Kestrel Estates LLP", "Entity in which a director is a designated partner"),
                       ("Kulkarni Resins LLP", "Entity in which a director is a partner")],
        "rp_txns": [("Rohan Kulkarni", "Managerial remuneration", 48.00, 42.00),
                    ("Sneha Kulkarni", "Rent for office premises", 9.60, 9.60),
                    ("Kestrel Estates LLP", "Sale of land and building", 480.00, 0.00),
                    ("Kulkarni Resins LLP", "Purchase of raw materials", 186.40, 172.10)],
        "ppe_note": (
            "During the year the Company sold its land and building at Plot 14, Sinnar Industrial "
            "Area to Kestrel Estates LLP for a consideration of Rs. 480.00 lakhs. The sale was "
            "approved by the Board of Directors at its meeting held on 14 August 2024."
        ),
        "aoc2": (
            "All contracts and arrangements with related parties during the year were in the "
            "ordinary course of business and on an arm's length basis. There were no material "
            "contracts or arrangements requiring the approval of members under Section 188(1) of "
            "the Act, and Form AOC-2 is therefore not applicable."
        ),
        "borrowings_note": (
            "Long-term borrowings are term loans from a scheduled bank secured by hypothecation of "
            "plant and machinery. Short-term borrowings are working capital facilities from the same bank."
        ),
        "extra_dr": [],
        "caro": False,
        "performance_note": "",
        "ledger": {"seed": 101, "n": 800, "rent": ("Sneha Kulkarni", 80000.00),
                   "security": ("Shield Guard Services", 60000.00),
                   "duplicates": 4, "large_payment": ("Legal and professional fees",
                                                      "Zenith Advisory Services", 18500000.00)},
        "issues": ["rpt_threshold_no_approval", "caro_missing", "duplicate_entries",
                   "unusual_large_payment"],
    },
    {
        "slug": "tarangini",
        "name": "Tarangini Foods Private Limited",
        "short": "Tarangini Foods",
        "city": "Mysuru", "state": "Karnataka", "incorporated": 2012,
        "business": "makes and sells packaged snacks and ready-to-eat savouries",
        "directors": [("Suresh Hegde", "Managing Director"), ("Anjali Hegde", "Director"),
                      ("Farhan Qureshi", "Director")],
        "auditor": ("Kaveri Assurance & Co.", "Dinesh Rao", "Bengaluru"),
        "report_date": "5 September 2025",
        "approval_limit": 50000,
        "pnl": {"revenue": (2964.80, 2718.35), "other_income": (18.40, 15.95),
                "materials": (1690.25, 1552.60), "employee": (402.15, 371.80),
                "finance": (64.30, 70.25), "depreciation": (96.70, 91.45),
                "other_exp": (438.60, 401.65)},
        "bs": {"share_capital": (250.00, 250.00), "other_equity_prior": 865.30,
               "lt_borrowings": (310.00, 355.00), "st_borrowings": (220.40, 198.60),
               "trade_payables": (356.85, 331.20), "other_cl": (98.40, 90.15),
               "ppe": (905.20, 872.40), "inventories": (402.60, 371.95),
               "receivables": (431.10, 396.40), "cash": (86.25, 79.80)},
        "other_exp_lines": [("Power and fuel", 0.22), ("Packing materials", None), ("Rent", None),
                            ("Repairs and maintenance", 0.09), ("Freight and forwarding", 0.16),
                            ("Advertisement and sales promotion", 0.08),
                            ("Travelling and conveyance", 0.04), ("Security charges", None)],
        "fixed_lines": {"Packing materials": (71.30, 64.85), "Rent": (7.80, 7.80),
                        "Security charges": (5.40, 5.10)},
        "footnotes": [
            "Packing materials include purchases of Rs. 61.84 lakhs from Nandi Print & Pack.",
        ],
        "mbp1": [("Suresh Hegde", "Hegde Agro Farms LLP", "Partner"),
                 ("Anjali Hegde", "Nandi Print & Pack", "Partner (50% share)"),
                 ("Farhan Qureshi", "Deccan Snack Distributors Private Limited", "Director")],
        "rp_parties": [("Suresh Hegde", "Managing Director (key managerial personnel)"),
                       ("Anjali Hegde", "Director; daughter of the Managing Director"),
                       ("Hegde Agro Farms LLP", "Entity in which a director is a partner")],
        "rp_txns": [("Suresh Hegde", "Managerial remuneration", 36.00, 33.00),
                    ("Anjali Hegde", "Remuneration", 18.00, 16.35),
                    ("Hegde Agro Farms LLP", "Purchase of raw materials", 118.40, 109.75)],
        "ppe_note": "Additions to plant and machinery during the year were Rs. 129.50 lakhs.",
        "aoc2": (
            "All related party transactions were in the ordinary course of business and on an arm's "
            "length basis and are disclosed in Note 27. Form AOC-2 is not applicable."
        ),
        "borrowings_note": (
            "Long-term borrowings are term loans from a scheduled bank secured by a charge on the "
            "factory land and building. Short-term borrowings are cash credit facilities."
        ),
        "extra_dr": [],
        "caro": True,
        "performance_note": "",
        "ledger": {"seed": 202, "n": 700, "rent": ("Chamundi Estates", 65000.00),
                   "security": ("Shield Guard Services", 45000.00),
                   "round_share": 0.35, "benford_plant": 120,
                   "rp_payments": ("Packing materials", "Nandi Print & Pack", 6184000.00, 14)},
        "issues": ["rp_omitted_from_schedule", "round_number_entries", "benford_nonconformity"],
    },
    {
        "slug": "northfield",
        "name": "Northfield Fabricators Private Limited",
        "short": "Northfield Fabricators",
        "city": "Hosur", "state": "Tamil Nadu", "incorporated": 2015,
        "business": "fabricates sheet-metal enclosures and brackets for electrical equipment makers",
        "directors": [("Prakash Iyer", "Managing Director"), ("Lakshmi Iyer", "Director")],
        "auditor": ("Nilgiri & Associates", "Karthik Subramanian", "Chennai"),
        "report_date": "12 September 2025",
        "approval_limit": 75000,
        "pnl": {"revenue": (846.10, 812.40), "other_income": (3.20, 2.85),
                "materials": (410.25, 580.30), "employee": (118.40, 112.75),
                "finance": (5.10, 5.85), "depreciation": (14.35, 13.90),
                "other_exp": (64.60, 61.20)},
        "bs": {"share_capital": (20.00, 20.00), "other_equity_prior": 186.40,
               "lt_borrowings": (42.00, 50.00), "st_borrowings": (28.50, 24.00),
               "trade_payables": (88.30, 92.15), "other_cl": (26.20, 24.60),
               "ppe": (148.60, 139.25), "inventories": (171.85, 96.40),
               "receivables": (142.30, 135.80), "cash": (21.40, 17.85)},
        "other_exp_lines": [("Power and fuel", 0.34), ("Rent", None), ("Repairs and maintenance", 0.14),
                            ("Freight and forwarding", 0.16), ("Travelling and conveyance", 0.06),
                            ("Security charges", None)],
        "fixed_lines": {"Rent": (4.56, 4.32), "Security charges": (2.88, 2.76)},
        "footnotes": [],
        "mbp1": [("Prakash Iyer", "Hosur Laser Cutting Works", "Proprietor")],
        "rp_parties": [("Prakash Iyer", "Managing Director (key managerial personnel)"),
                       ("Lakshmi Iyer", "Director; wife of the Managing Director"),
                       ("Venkat Raman", "Brother of Mrs. Lakshmi Iyer (Director)")],
        "rp_txns": [("Prakash Iyer", "Managerial remuneration", 24.00, 24.00),
                    ("Lakshmi Iyer", "Remuneration", 12.00, 12.00),
                    ("Venkat Raman", "Unsecured loan accepted", 4.50, 0.00)],
        "ppe_note": "Additions to plant and machinery during the year were Rs. 23.70 lakhs.",
        "aoc2": (
            "All related party transactions were in the ordinary course of business and on an arm's "
            "length basis. Form AOC-2 is not applicable."
        ),
        "borrowings_note": (
            "Long-term borrowings are a term loan from a scheduled bank. Short-term borrowings include "
            "an unsecured loan of Rs. 4.50 lakhs accepted in cash on 18 November 2024 from a relative "
            "of a director (see Note 27); the loan is interest-free and repayable on demand."
        ),
        "extra_dr": [],
        "caro": True,
        "performance_note": (
            "Profit improved on account of better material yields during the year."
        ),
        "ledger": {"seed": 303, "n": 800, "rent": ("Hosur Industrial Sheds", 38000.00),
                   "security": ("Metro Watch Security", 24000.00), "duplicates": 5},
        "issues": ["duplicate_entries", "ratio_anomaly", "cash_loan_269ss"],
    },
    {
        "slug": "vardhan",
        "name": "Vardhan Precision Castings Private Limited",
        "short": "Vardhan Precision Castings",
        "city": "Rajkot", "state": "Gujarat", "incorporated": 2004,
        "business": "produces investment castings for pumps, valves and auto components",
        "directors": [("Harsh Vardhan Shah", "Managing Director"), ("Nisha Shah", "Director"),
                      ("Imran Patel", "Director")],
        "auditor": ("Tapi Assurance LLP", "Jignesh Mehta", "Ahmedabad"),
        "report_date": "22 August 2025",
        "approval_limit": 100000,
        "pnl": {"revenue": (5124.60, 4702.35), "other_income": (51.30, 47.10),
                "materials": (3105.40, 2842.75), "employee": (688.20, 640.35),
                "finance": (142.75, 151.20), "depreciation": (236.40, 221.85),
                "other_exp": (521.65, 482.95)},
        "bs": {"share_capital": (600.00, 600.00), "other_equity_prior": 2140.60,
               "lt_borrowings": (850.00, 920.00), "st_borrowings": (410.20, 385.40),
               "trade_payables": (620.45, 571.30), "other_cl": (188.30, 176.45),
               "ppe": (2440.80, 2105.60), "inventories": (702.35, 648.90),
               "receivables": (915.60, 842.25), "cash": (204.10, 190.35)},
        "other_exp_lines": [("Power and fuel", 0.36), ("Rent", None), ("Repairs and maintenance", 0.12),
                            ("Freight and forwarding", 0.13), ("Legal and professional fees", 0.04),
                            ("Travelling and conveyance", 0.04), ("Insurance", 0.03),
                            ("Security charges", None)],
        "fixed_lines": {"Rent": (13.32, 12.72), "Security charges": (8.40, 7.80)},
        "footnotes": [
            "Repairs and maintenance include Rs. 3.24 lakhs paid to Vardhan Engineering Works (see Note 27).",
        ],
        "mbp1": [("Harsh Vardhan Shah", "Vardhan Engineering Works", "Partner"),
                 ("Nisha Shah", "Shah Alloys LLP", "Designated partner"),
                 ("Imran Patel", "Saurashtra Pumps Private Limited", "Director")],
        "rp_parties": [("Harsh Vardhan Shah", "Managing Director (key managerial personnel)"),
                       ("Nisha Shah", "Director; wife of the Managing Director"),
                       ("Vardhan Engineering Works", "Firm in which a director is a partner"),
                       ("Shah Alloys LLP", "Entity in which a director is a designated partner")],
        "rp_txns": [("Harsh Vardhan Shah", "Managerial remuneration", 60.00, 54.00),
                    ("Nisha Shah", "Remuneration", 24.00, 21.00),
                    ("Vardhan Engineering Works", "Purchase of machinery", 352.80, 0.00),
                    ("Vardhan Engineering Works", "Repairs and maintenance services", 3.24, 2.80),
                    ("Shah Alloys LLP", "Purchase of raw materials", 214.35, 198.10)],
        "ppe_note": (
            "Additions to plant and machinery include a vacuum casting line purchased from Vardhan "
            "Engineering Works for Rs. 352.80 lakhs."
        ),
        "aoc2": (
            "The purchase of a vacuum casting line from Vardhan Engineering Works for Rs. 352.80 lakhs "
            "exceeded the threshold in Rule 15 of the Companies (Meetings of Board and its Powers) "
            "Rules, 2014. It was approved by the Board on 30 July 2024 and by the members by special "
            "resolution at the extraordinary general meeting held on 20 August 2024, and the "
            "particulars are given in Form AOC-2 annexed to this report. All other related party "
            "transactions were in the ordinary course of business and on an arm's length basis."
        ),
        "borrowings_note": (
            "Long-term borrowings are term loans from scheduled banks secured by plant and machinery. "
            "Short-term borrowings are cash credit facilities secured by inventories and receivables."
        ),
        "extra_dr": [
            "During the year the Company did not accept any loan or deposit in cash; all borrowings "
            "were received through banking channels.",
        ],
        "caro": True,
        "performance_note": "",
        "ledger": {"seed": 404, "n": 800, "rent": ("Bhakti Industrial Sheds", 111000.00),
                   "security": ("Shield Guard Services", 70000.00),
                   "rp_payments": ("Repairs and maintenance", "Vardhan Engineering Works", 324000.00, 3)},
        "issues": [],
        "decoys": [
            "Material related-party purchase (Rs. 352.80 lakhs, above 10% of net worth) that WAS "
            "approved by special resolution and disclosed in Form AOC-2",
            "A negative statement: no loan or deposit was accepted in cash",
            "Director-linked vendor (Vardhan Engineering Works) that appears in the ledger AND in the "
            "related-party note",
        ],
    },
    {
        "slug": "meridian",
        "name": "Meridian Agro Exports Private Limited",
        "short": "Meridian Agro Exports",
        "city": "Kochi", "state": "Kerala", "incorporated": 2011,
        "business": "processes and exports whole and ground spices",
        "directors": [("Thomas Varghese", "Managing Director"), ("Maria Varghese", "Director"),
                      ("Ravi Menon", "Director")],
        "auditor": ("Vembanad Audit & Co.", "Anitha Pillai", "Kochi"),
        "report_date": "15 September 2025",
        "approval_limit": 100000,
        "pnl": {"revenue": (3648.90, 3442.35), "other_income": (64.20, 58.75),
                "materials": (2446.30, 2340.80), "employee": (318.45, 296.20),
                "finance": (96.80, 88.45), "depreciation": (74.25, 70.10),
                "other_exp": (402.85, 381.75)},
        "bs": {"share_capital": (300.00, 300.00), "other_equity_prior": 1024.50,
               "lt_borrowings": (280.00, 310.00), "st_borrowings": (640.25, 402.10),
               "trade_payables": (368.40, 351.75), "other_cl": (112.60, 104.30),
               "ppe": (612.40, 596.85), "inventories": (546.20, 521.35),
               "receivables": (1102.80, 565.40), "cash": (58.90, 71.25)},
        "other_exp_lines": [("Power and fuel", 0.18), ("Cold storage charges", None), ("Rent", None),
                            ("Repairs and maintenance", 0.07), ("Freight and forwarding", 0.26),
                            ("Export documentation charges", 0.05),
                            ("Travelling and conveyance", 0.05), ("Security charges", None)],
        "fixed_lines": {"Cold storage charges": (49.60, 45.10), "Rent": (14.64, 14.64),
                        "Security charges": (6.00, 5.70)},
        "footnotes": [
            "Cold storage charges include Rs. 42.18 lakhs paid to Periyar Cold Storage LLP.",
            "Rent is paid to Varghese Holdings Private Limited, the holding company (see Note 27).",
        ],
        "mbp1": [("Thomas Varghese", "Varghese Holdings Private Limited", "Director and shareholder"),
                 ("Maria Varghese", "Periyar Cold Storage LLP", "Designated partner"),
                 ("Ravi Menon", "Coastal Freight Forwarders Private Limited", "Director")],
        "rp_parties": [("Varghese Holdings Private Limited", "Holding company (holds 74% of equity)"),
                       ("Thomas Varghese", "Managing Director (key managerial personnel)"),
                       ("Maria Varghese", "Director; wife of the Managing Director")],
        "rp_txns": [("Thomas Varghese", "Managerial remuneration", 42.00, 39.00),
                    ("Maria Varghese", "Remuneration", 18.00, 18.00),
                    ("Varghese Holdings Private Limited", "Rent for warehouse", 14.64, 14.64)],
        "ppe_note": "Additions to plant and machinery during the year were Rs. 89.80 lakhs.",
        "aoc2": (
            "All related party transactions were in the ordinary course of business and on an arm's "
            "length basis and are disclosed in Note 27. Form AOC-2 is not applicable."
        ),
        "borrowings_note": (
            "Long-term borrowings are a term loan from a scheduled bank. Short-term borrowings are "
            "packing credit and export bill discounting facilities."
        ),
        "extra_dr": [],
        "caro": False,
        "performance_note": "",
        "ledger": {"seed": 505, "n": 800, "rent": ("Varghese Holdings Private Limited", 122000.00),
                   "security": ("Metro Watch Security", 50000.00), "round_share": 0.35,
                   "rp_payments": ("Cold storage charges", "Periyar Cold Storage LLP", 4218000.00, 12)},
        "issues": ["rp_omitted_from_schedule", "caro_missing", "round_number_entries", "ratio_anomaly"],
    },
]

# Ledger accounts and vendors shared by all companies (fictional names)
LEDGER_ACCOUNTS = {
    "Power and fuel": ["Suryoday Power Services", "Urja Diesel Suppliers", "GreenGrid Energy Services"],
    "Repairs and maintenance": ["Apex Tools & Spares", "Precision Machine Care", "Unity Electricals",
                                "Sai Civil Works"],
    "Freight and forwarding": ["Swift Roadlines", "Ganga Freight Carriers", "Coastline Cargo Movers"],
    "Travelling and conveyance": ["Skyway Travels", "City Cab Services", "Hotel Residency Inn"],
    "Printing and stationery": ["Metro Stationers", "Quickprint Solutions"],
    "Communication expenses": ["Bluewave Communications", "NetLink Broadband"],
    "Insurance": ["Pioneer Insurance Brokers"],
    "Legal and professional fees": ["Kapoor & Sen Legal Associates", "Ledger Tax Consultants",
                                    "Brightpath HR Advisory"],
    "Housekeeping": ["Clean Sweep Facility Services"],
    "Miscellaneous expenses": ["Sundry Vendors", "Office Mart", "Canteen Services Co-op"],
}


# ---------------------------------------------------------------------------
# Statement computation
# ---------------------------------------------------------------------------
def r2(x: float) -> float:
    return round(x + 1e-9, 2)


def compute_statements(c: dict) -> dict:
    """Derive totals, tax, profit, other equity and the balancing asset line."""
    out: dict = {"pnl": {}, "bs": {}}
    for i, _ in enumerate((CUR, PRI)):
        p = {k: c["pnl"][k][i] for k, _ in PNL_KEYS}
        total_income = r2(p["revenue"] + p["other_income"])
        total_exp = r2(p["materials"] + p["employee"] + p["finance"] + p["depreciation"] + p["other_exp"])
        pbt = r2(total_income - total_exp)
        tax = r2(pbt * TAX_RATE)
        pat = r2(pbt - tax)
        p.update(total_income=total_income, total_expenses=total_exp, pbt=pbt, tax=tax, pat=pat)
        out["pnl"][i] = p
    oe_prior = c["bs"]["other_equity_prior"]
    oe_cur = r2(oe_prior + out["pnl"][0]["pat"])
    for i, oe in ((0, oe_cur), (1, oe_prior)):
        b = {k: c["bs"][k][i] for k, _ in BS_KEYS}
        b["other_equity"] = oe
        total_el = r2(b["share_capital"] + oe + b["lt_borrowings"] + b["st_borrowings"]
                      + b["trade_payables"] + b["other_cl"])
        known_assets = r2(b["ppe"] + b["inventories"] + b["receivables"] + b["cash"])
        b["other_ca"] = r2(total_el - known_assets)
        if b["other_ca"] <= 0:
            raise ValueError(f"{c['slug']}: balancing other current assets is {b['other_ca']}")
        b["total_el"] = total_el
        b["total_assets"] = total_el
        b["net_worth"] = r2(b["share_capital"] + oe)
        out["bs"][i] = b
    # Other expenses breakdown: fixed lines, share-based lines, balancing miscellaneous line
    lines = []
    for i in (0, 1):
        total = c["pnl"]["other_exp"][i]
        vals = {}
        for name, share in c["other_exp_lines"]:
            vals[name] = c["fixed_lines"][name][i] if share is None else r2(total * share)
        vals["Miscellaneous expenses"] = r2(total - sum(vals.values()))
        if vals["Miscellaneous expenses"] <= 0:
            raise ValueError(f"{c['slug']}: other expenses breakdown exceeds the total")
        lines.append(vals)
    out["other_exp_breakdown"] = [(k, lines[0][k], lines[1][k]) for k in lines[0]]
    return out


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------
def fmt(x: float) -> str:
    return "–" if x == 0 else f"{x:,.2f}"


def md_table(header: list[str], rows: list[list[str]], right_from: int = 1) -> str:
    align = ["---" if i < right_from else "---:" for i in range(len(header))]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(align) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in rows]
    return "\n".join(lines)


def pct(a: float, b: float) -> float:
    return (a - b) / b * 100


# ---------------------------------------------------------------------------
# Annual report (Markdown)
# ---------------------------------------------------------------------------
def pnl_rows(st: dict) -> list[tuple[str, str, float, float]]:
    p0, p1 = st["pnl"][0], st["pnl"][1]
    return [
        ("Revenue from operations", "20", p0["revenue"], p1["revenue"]),
        ("Other income", "21", p0["other_income"], p1["other_income"]),
        ("Total income", "", p0["total_income"], p1["total_income"]),
        ("Cost of materials consumed", "22", p0["materials"], p1["materials"]),
        ("Employee benefits expense", "23", p0["employee"], p1["employee"]),
        ("Finance costs", "24", p0["finance"], p1["finance"]),
        ("Depreciation and amortisation expense", "3", p0["depreciation"], p1["depreciation"]),
        ("Other expenses", "25", p0["other_exp"], p1["other_exp"]),
        ("Total expenses", "", p0["total_expenses"], p1["total_expenses"]),
        ("Profit before tax", "", p0["pbt"], p1["pbt"]),
        ("Tax expense", "26", p0["tax"], p1["tax"]),
        ("Profit for the year", "", p0["pat"], p1["pat"]),
    ]


def bs_rows(st: dict) -> list[tuple[str, str, float, float]]:
    b0, b1 = st["bs"][0], st["bs"][1]
    return [
        ("Equity share capital", "4", b0["share_capital"], b1["share_capital"]),
        ("Other equity", "5", b0["other_equity"], b1["other_equity"]),
        ("Long-term borrowings", "6", b0["lt_borrowings"], b1["lt_borrowings"]),
        ("Short-term borrowings", "6", b0["st_borrowings"], b1["st_borrowings"]),
        ("Trade payables", "7", b0["trade_payables"], b1["trade_payables"]),
        ("Other current liabilities", "8", b0["other_cl"], b1["other_cl"]),
        ("Total equity and liabilities", "", b0["total_el"], b1["total_el"]),
        ("Property, plant and equipment", "3", b0["ppe"], b1["ppe"]),
        ("Inventories", "9", b0["inventories"], b1["inventories"]),
        ("Trade receivables", "10", b0["receivables"], b1["receivables"]),
        ("Cash and cash equivalents", "11", b0["cash"], b1["cash"]),
        ("Other current assets", "12", b0["other_ca"], b1["other_ca"]),
        ("Total assets", "", b0["total_assets"], b1["total_assets"]),
    ]


def build_markdown(c: dict, st: dict) -> str:
    name, short = c["name"], c["short"]
    p0, p1 = st["pnl"][0], st["pnl"][1]
    firm, partner, place = c["auditor"]
    directors = ", ".join(f"{n} ({d})" for n, d in c["directors"])
    md: list[str] = []
    md.append(f"# {name}\n\nAnnual report for the financial year 2024-25 (extract)\n\n*{DISCLAIMER}*")

    md.append("## Corporate information\n\n"
              f"{name} (\"the Company\") is a private limited company incorporated in {c['incorporated']} "
              f"with its registered office at {c['city']}, {c['state']}. Its shares are not listed on any "
              f"stock exchange. The Company {c['business']}. The directors of the Company are {directors}.")

    perf = (f"Revenue from operations for the year was Rs. {fmt(p0['revenue'])} lakhs against "
            f"Rs. {fmt(p1['revenue'])} lakhs in the previous year, and the profit for the year was "
            f"Rs. {fmt(p0['pat'])} lakhs against Rs. {fmt(p1['pat'])} lakhs.")
    if c["performance_note"]:
        perf += " " + c["performance_note"]
    dr = [f"## Directors' report (extract) for the year ended 31 March 2025\n\n"
          f"### Financial performance\n\n{perf} The Board does not recommend a dividend.",
          "### Disclosure of interest by directors (Form MBP-1)\n\n"
          "The directors have disclosed the following interests in other entities:\n\n"
          + md_table(["Director", "Entity", "Nature of interest"],
                     [[d, e, n] for d, e, n in c["mbp1"]], right_from=3),
          "### Particulars of contracts or arrangements with related parties (Section 188)\n\n" + c["aoc2"],
          "### Internal financial controls\n\n"
          "The Company has documented policies for purchases and payments. Any single payment above "
          f"Rs. {c['approval_limit']:,} requires the approval of two directors. The Board considers the "
          "internal financial controls to be adequate and operating effectively.",
          "### Deposits\n\nThe Company has not accepted any deposits from the public within the meaning "
          "of Chapter V of the Companies Act, 2013." + ("".join(" " + x for x in c["extra_dr"]))]
    md.extend(dr)

    md.append("## Balance Sheet as at 31 March 2025\n\n(All amounts in Rs. lakhs)\n\n"
              + md_table(["Particulars", "Note", "31 Mar 2025", "31 Mar 2024"],
                         [[a, n, fmt(x), fmt(y)] for a, n, x, y in bs_rows(st)], right_from=2))
    md.append("## Statement of Profit and Loss for the year ended 31 March 2025\n\n(All amounts in Rs. lakhs)\n\n"
              + md_table(["Particulars", "Note", "FY 2024-25", "FY 2023-24"],
                         [[a, n, fmt(x), fmt(y)] for a, n, x, y in pnl_rows(st)], right_from=2))

    notes = ["## Notes to the financial statements (selected)",
             "### Note 1: Significant accounting policies\n\n"
             "The financial statements are prepared on a going concern basis under the historical cost "
             "convention and the accrual basis of accounting, and comply with the Accounting Standards "
             "notified under Section 133 of the Companies Act, 2013. All amounts are in Rs. lakhs unless "
             "stated otherwise.",
             f"### Note 3: Property, plant and equipment\n\n{c['ppe_note']}",
             f"### Note 6: Borrowings\n\n{c['borrowings_note']}",
             "### Note 25: Other expenses\n\n(All amounts in Rs. lakhs)\n\n"
             + md_table(["Particulars", "FY 2024-25", "FY 2023-24"],
                        [[k, fmt(x), fmt(y)] for k, x, y in st["other_exp_breakdown"]]
                        + [["Total", fmt(p0["other_exp"]), fmt(p1["other_exp"])]])
             + "".join(f"\n\n{i + 1}. {f}" for i, f in enumerate(c["footnotes"]))]
    rp_list = "\n".join(f"- {p}: {rel}" for p, rel in c["rp_parties"])
    rel_of = dict(c["rp_parties"])
    rp_rows = [[p, rel_of[p].split(";")[0], nat, fmt(x), fmt(y)] for p, nat, x, y in c["rp_txns"]]
    notes.append("### Note 27: Related party disclosures (AS 18)\n\n"
                 f"Names of related parties and nature of relationship:\n\n{rp_list}\n\n"
                 "Transactions with related parties (All amounts in Rs. lakhs):\n\n"
                 + md_table(["Related party", "Relationship", "Nature of transaction", "FY 2024-25", "FY 2023-24"],
                            rp_rows, right_from=3))
    notes.append("### Note 28: Contingent liabilities\n\nClaims against the Company not acknowledged as "
                 "debts: Nil (previous year: Nil).")
    md.extend(notes)

    # Independent auditor's report
    ar = [f"## Independent Auditor's Report\n\nTo the Members of {name}",
          "### Report on the audit of the financial statements\n\n"
          f"**Opinion.** We have audited the financial statements of {name} (\"the Company\"), which "
          "comprise the Balance Sheet as at 31 March 2025, the Statement of Profit and Loss for the year "
          "then ended and notes to the financial statements, including a summary of significant "
          "accounting policies. In our opinion and to the best of our information and according to the "
          "explanations given to us, the financial statements give the information required by the "
          "Companies Act, 2013 (\"the Act\") in the manner so required and give a true and fair view in "
          "conformity with the accounting principles generally accepted in India of the state of affairs "
          "of the Company as at 31 March 2025 and its profit for the year ended on that date.\n\n"
          "**Basis for opinion.** We conducted our audit in accordance with the Standards on Auditing "
          "specified under Section 143(10) of the Act. We are independent of the Company in accordance "
          "with the Code of Ethics issued by the Institute of Chartered Accountants of India, and we "
          "believe that the audit evidence we have obtained is sufficient and appropriate to provide a "
          "basis for our opinion."]
    legal = ["### Report on other legal and regulatory requirements"]
    items = []
    if c["caro"]:
        items.append("As required by the Companies (Auditor's Report) Order, 2020 (\"the Order\"), issued by "
                     "the Central Government in terms of Section 143(11) of the Act, we give in Annexure A a "
                     "statement on the matters specified in paragraphs 3 and 4 of the Order, to the extent "
                     "applicable.")
    items.append("As required by Section 143(3) of the Act, we report that: (a) we have sought and obtained "
                 "all the information and explanations which to the best of our knowledge and belief were "
                 "necessary for the purposes of our audit; (b) in our opinion, proper books of account as "
                 "required by law have been kept by the Company; (c) the Balance Sheet and the Statement of "
                 "Profit and Loss dealt with by this report are in agreement with the books of account; "
                 "(d) in our opinion, the financial statements comply with the Accounting Standards "
                 "specified under Section 133 of the Act; (e) none of the directors is disqualified under "
                 "Section 164(2) of the Act; (f) the Company has adequate internal financial controls with "
                 "reference to financial statements and they were operating effectively; (g) the Company "
                 "has no pending litigations which would impact its financial position.")
    legal.append("\n\n".join(f"{i + 1}. {t}" for i, t in enumerate(items)))
    legal.append(f"For {firm}, Chartered Accountants\n\n{partner}, Partner\n\nPlace: {place}\n\n"
                 f"Date: {c['report_date']}")
    ar.extend(legal)
    if c["caro"]:
        ar.append("### Annexure A to the Independent Auditor's Report (Companies (Auditor's Report) Order, 2020)\n\n"
                  "(i) The Company maintains proper records of property, plant and equipment, which were "
                  "physically verified by management during the year.\n\n"
                  "(ii) Inventories were physically verified by management at reasonable intervals and no "
                  "material discrepancies were noticed.\n\n"
                  "(iii) The Company has not granted any loans or advances to companies, firms or other "
                  "parties.\n\n"
                  "(v) The Company has not accepted deposits from the public.\n\n"
                  "(vii) The Company has been regular in depositing undisputed statutory dues, including "
                  "goods and services tax, provident fund and income tax, with the appropriate authorities.\n\n"
                  "(ix) The Company has not defaulted in the repayment of loans or borrowings to any lender.\n\n"
                  "(xi) No fraud by the Company or on the Company has been noticed or reported during the year.\n\n"
                  "(xiii) Transactions with related parties are in compliance with Section 188 of the Act "
                  "where applicable, and the details have been disclosed in the financial statements as "
                  "required by the applicable Accounting Standards.")
    md.extend(ar)
    return "\n\n".join(md) + "\n"


# ---------------------------------------------------------------------------
# Ledger
# ---------------------------------------------------------------------------
def _random_date(rng: random.Random) -> dt.date:
    return FY_START + dt.timedelta(days=rng.randrange((FY_END - FY_START).days + 1))


def _split_amount(rng: random.Random, total: float, n: int) -> list[float]:
    """Split a rupee total into n positive two-decimal amounts that add up exactly."""
    weights = [rng.uniform(0.5, 1.5) for _ in range(n)]
    s = sum(weights)
    paise = round(total * 100)
    parts = [int(paise * w / s) for w in weights]
    parts[-1] += paise - sum(parts)
    return [p / 100 for p in parts]


def build_ledger(c: dict) -> tuple[list[dict], dict]:
    """Return ledger rows and a record of what was planted where."""
    cfg = c["ledger"]
    rng = random.Random(cfg["seed"])
    accounts = list(LEDGER_ACCOUNTS)
    rows: list[dict] = []
    for _ in range(cfg["n"]):
        acct = rng.choice(accounts)
        rows.append({"date": _random_date(rng), "account": acct,
                     "party": rng.choice(LEDGER_ACCOUNTS[acct]),
                     # log-uniform over exactly four decades (Rs. 100 to Rs. 10 lakh), so the
                     # leading digits follow Benford's law
                     "amount": round(10 ** rng.uniform(2, 6), 2)})
    planted: dict = {}

    # Round-number plant: round a share of the entries (those of Rs. 2,000 or more) to the
    # nearest thousand
    if cfg.get("round_share"):
        eligible = [i for i, r in enumerate(rows) if r["amount"] >= 2000]
        idx = rng.sample(eligible, int(len(rows) * cfg["round_share"]))
        for i in idx:
            rows[i]["amount"] = float(round(rows[i]["amount"] / 1000) * 1000)
        planted["round_number_rows"] = len(idx)

    # Benford plant: invoices just below the approval limit
    if cfg.get("benford_plant"):
        lim = c["approval_limit"]
        for _ in range(cfg["benford_plant"]):
            acct = rng.choice(["Repairs and maintenance", "Miscellaneous expenses"])
            rows.append({"date": _random_date(rng), "account": acct,
                         "party": rng.choice(LEDGER_ACCOUNTS[acct]),
                         "amount": round(rng.uniform(0.8 * lim, lim - 0.01), 2)})
        planted["just_below_limit_rows"] = cfg["benford_plant"]

    # Recurring monthly payments (legitimately identical amounts)
    for acct, (party, amount) in (("Rent", cfg["rent"]), ("Security charges", cfg["security"])):
        for m in range(12):
            month = (FY_START.month + m - 1) % 12 + 1
            year = FY_START.year + (FY_START.month + m - 1) // 12
            rows.append({"date": dt.date(year, month, 5), "account": acct, "party": party, "amount": amount})

    # Payments to a director-linked vendor
    if cfg.get("rp_payments"):
        acct, party, total, n = cfg["rp_payments"]
        for amt in _split_amount(rng, total, n):
            rows.append({"date": _random_date(rng), "account": acct, "party": party, "amount": amt})

    # One very large payment
    if cfg.get("large_payment"):
        acct, party, amount = cfg["large_payment"]
        rows.append({"date": dt.date(2025, 2, 17), "account": acct, "party": party, "amount": amount})
        planted["large_payment"] = amount

    rows.sort(key=lambda r: (r["date"], r["account"], r["party"], r["amount"]))

    # Duplicate plant: the same voucher posted twice (same date, account, party, amount)
    if cfg.get("duplicates"):
        base = [r for r in rows if r["account"] not in ("Rent", "Security charges")]
        dups = rng.sample(base, cfg["duplicates"])
        rows.extend(dict(r) for r in dups)
        rows.sort(key=lambda r: (r["date"], r["account"], r["party"], r["amount"]))
        planted["duplicate_pairs"] = [
            {"date": r["date"].isoformat(), "account": r["account"], "party": r["party"],
             "amount": r["amount"]} for r in dups]

    for i, r in enumerate(rows, start=1):
        r["voucher"] = f"PV/24-25/{i:04d}"
    return rows, planted


def write_ledger_csv(rows: list[dict], path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Date", "Voucher No", "Account", "Party", "Amount"])
        for r in rows:
            w.writerow([r["date"].isoformat(), r["voucher"], r["account"], r["party"], f"{r['amount']:.2f}"])


# ---------------------------------------------------------------------------
# Workbook and PDF (byte-stable output)
# ---------------------------------------------------------------------------
_FIXED_TIME = dt.datetime(2025, 9, 30, 0, 0, 0)


def _restamp_zip(path: Path) -> None:
    """Rewrite a zip container with fixed timestamps so the bytes are reproducible."""
    import re

    with zipfile.ZipFile(path) as zin:
        items = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    stamp = _FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    # openpyxl sets the "modified" property to the save time; pin it
    items = [(n, re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)", rb"\g<1>" + stamp + rb"\2", d)
              if n == "docProps/core.xml" else d) for n, d in items]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items:
            info = zipfile.ZipInfo(name, date_time=_FIXED_TIME.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(info, data)
    path.write_bytes(buf.getvalue())


def write_xlsx(c: dict, st: dict, path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"
    ws.append([c["name"]])
    ws.append(["Balance Sheet as at 31 March 2025"])
    ws.append(["(All amounts in Rs. lakhs)"])
    ws.append([])
    ws.append(["Particulars", BS_CUR, BS_PRI])
    for a, _, x, y in bs_rows(st):
        ws.append([a, x, y])

    ws = wb.create_sheet("Profit and Loss")
    ws.append([c["name"]])
    ws.append(["Statement of Profit and Loss for the year ended 31 March 2025"])
    ws.append(["(All amounts in Rs. lakhs)"])
    ws.append([])
    ws.append(["Particulars", CUR, PRI])
    for a, _, x, y in pnl_rows(st):
        ws.append([a, x, y])

    ws = wb.create_sheet("Related Party Transactions")
    ws.append([c["name"]])
    ws.append(["Related party transactions (All amounts in Rs. lakhs)"])
    ws.append([])
    ws.append(["Name of Related Party", "Relationship", "Nature of Transaction", CUR, PRI])
    rel_of = dict(c["rp_parties"])
    for p, nat, x, y in c["rp_txns"]:
        ws.append([p, rel_of[p].split(";")[0], nat, x, y])

    wb.properties.creator = "benchmark generator"
    wb.properties.created = _FIXED_TIME
    wb.properties.modified = _FIXED_TIME
    wb.save(path)
    _restamp_zip(path)


def _md_to_flowables(md: str):
    """Render the small Markdown subset used by build_markdown with reportlab."""
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    small = styles["BodyText"].clone("small", fontSize=8, leading=10)
    out = []

    def inline(text: str) -> str:
        text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        while "**" in text:
            text = text.replace("**", "<b>", 1).replace("**", "</b>", 1)
        if text.startswith("*") and text.endswith("*"):
            text = f"<i>{text[1:-1]}</i>"
        return text

    for block in md.strip().split("\n\n"):
        lines = block.split("\n")
        if block.startswith("#"):
            level = len(block) - len(block.lstrip("#"))
            style = {1: "Title", 2: "Heading2", 3: "Heading3"}.get(level, "Heading4")
            out.append(Paragraph(inline(block.lstrip("#").strip()), styles[style]))
        elif lines[0].startswith("|"):
            cells = [[c.strip() for c in ln.strip("|").split("|")] for ln in lines if not set(ln) <= set("|-: ")]
            data = [[Paragraph(inline(c), small) for c in row] for row in cells]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            out.append(t)
        else:
            for ln in lines:
                out.append(Paragraph(inline(ln), body))
        out.append(Spacer(1, 4))
    return out


def write_pdf(c: dict, md: str, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate

    doc = SimpleDocTemplate(str(path), pagesize=A4, title=f"{c['name']} - Annual report FY 2024-25 (fictional)",
                            author="benchmark generator", subject=DISCLAIMER, creator="benchmark generator",
                            invariant=1)
    doc.build(_md_to_flowables(md))


# ---------------------------------------------------------------------------
# Questions and ground truth
# ---------------------------------------------------------------------------
def build_questions(c: dict, st: dict) -> list[dict]:
    p0, p1, b0, b1 = st["pnl"][0], st["pnl"][1], st["bs"][0], st["bs"][1]
    s, n = c["slug"], c["name"]
    rp = {(p, nat): x for p, nat, x, _ in c["rp_txns"]}

    def num(qid, cat, q, gold, unit="lakhs"):
        return {"id": f"{s}-{qid}", "company": s, "category": cat, "question": q,
                "answer_type": "number", "unit": unit, "gold": round(gold, 2)}

    def text(qid, cat, q, gold, aliases=()):
        return {"id": f"{s}-{qid}", "company": s, "category": cat, "question": q,
                "answer_type": "text", "gold": gold, "aliases": [gold, *aliases]}

    def date(qid, cat, q, gold):
        return {"id": f"{s}-{qid}", "company": s, "category": cat, "question": q,
                "answer_type": "date", "gold": gold}

    firm = c["auditor"][0]
    if s == "kestrel":
        return [
            num("q1", "lookup", f"What was {n}'s revenue from operations for FY 2024-25, in Rs. lakhs?", p0["revenue"]),
            num("q2", "calculation", f"By how many Rs. lakhs did {n}'s employee benefits expense increase from FY 2023-24 to FY 2024-25?",
                p0["employee"] - p1["employee"]),
            num("q3", "related_party", f"For what consideration, in Rs. lakhs, did {n} sell land and building to a related party in FY 2024-25?",
                rp[("Kestrel Estates LLP", "Sale of land and building")]),
            text("q4", "related_party", f"Which related party bought land and building from {n} during FY 2024-25?",
                 "Kestrel Estates LLP", ["Kestrel Estates"]),
            text("q5", "multi_hop", f"The related party that bought land and building from {n} in FY 2024-25 is connected to which director of the company?",
                 "Rohan Kulkarni"),
            text("q6", "lookup", f"Which firm is the statutory auditor of {n}?", firm, ["Sahyadri Audit Partners"]),
        ]
    if s == "tarangini":
        return [
            num("q1", "lookup", f"What was {n}'s profit for the year for FY 2024-25, in Rs. lakhs?", p0["pat"]),
            num("q2", "calculation", f"By what percentage did {n}'s revenue from operations grow in FY 2024-25 compared with FY 2023-24?",
                pct(p0["revenue"], p1["revenue"]), unit="percent"),
            num("q3", "related_party", f"What value of raw materials, in Rs. lakhs, did {n} purchase from Hegde Agro Farms LLP in FY 2024-25?",
                rp[("Hegde Agro Farms LLP", "Purchase of raw materials")]),
            text("q4", "multi_hop", f"{n} bought packing materials from a firm in which one of its directors is a partner. Which director?",
                 "Anjali Hegde"),
            num("q5", "multi_hop", f"What value of packing materials, in Rs. lakhs, did {n} purchase in FY 2024-25 from the firm in which its director Anjali Hegde is a partner?",
                61.84),
            num("q6", "lookup", f"What were {n}'s trade receivables as at 31 March 2025, in Rs. lakhs?", b0["receivables"]),
        ]
    if s == "northfield":
        return [
            num("q1", "lookup", f"What was {n}'s cost of materials consumed in FY 2024-25, in Rs. lakhs?", p0["materials"]),
            num("q2", "calculation", f"What was {n}'s gross margin in FY 2023-24, defined as revenue from operations minus cost of materials consumed, as a percentage of revenue from operations?",
                (p1["revenue"] - p1["materials"]) / p1["revenue"] * 100, unit="percent"),
            num("q3", "calculation", f"By how many Rs. lakhs did {n}'s inventories increase between 31 March 2024 and 31 March 2025?",
                b0["inventories"] - b1["inventories"]),
            text("q4", "multi_hop", f"From whom did {n} accept a loan in cash during FY 2024-25?", "Venkat Raman"),
            num("q5", "related_party", f"What was the amount, in Rs. lakhs, of the unsecured loan {n} accepted in cash during FY 2024-25?",
                rp[("Venkat Raman", "Unsecured loan accepted")]),
            text("q6", "lookup", f"Who is the Managing Director of {n}?", "Prakash Iyer"),
        ]
    if s == "vardhan":
        return [
            num("q1", "lookup", f"What were {n}'s finance costs in FY 2023-24, in Rs. lakhs?", p1["finance"]),
            num("q2", "calculation", f"By how many Rs. lakhs did {n}'s revenue from operations increase in FY 2024-25 compared with FY 2023-24?",
                p0["revenue"] - p1["revenue"]),
            num("q3", "related_party", f"What was the value, in Rs. lakhs, of the machinery {n} purchased from Vardhan Engineering Works in FY 2024-25?",
                rp[("Vardhan Engineering Works", "Purchase of machinery")]),
            date("q4", "related_party", f"On what date did the members of {n} approve the purchase of machinery from Vardhan Engineering Works?",
                 "2024-08-20"),
            text("q5", "multi_hop", f"Vardhan Engineering Works is a related party of {n} through which director?",
                 "Harsh Vardhan Shah"),
            num("q6", "lookup", f"What was {n}'s property, plant and equipment as at 31 March 2025, in Rs. lakhs?", b0["ppe"]),
        ]
    if s == "meridian":
        return [
            num("q1", "lookup", f"What were {n}'s trade receivables as at 31 March 2025, in Rs. lakhs?", b0["receivables"]),
            num("q2", "calculation", f"By what percentage did {n}'s trade receivables increase between 31 March 2024 and 31 March 2025?",
                pct(b0["receivables"], b1["receivables"]), unit="percent"),
            num("q3", "related_party", f"What rent, in Rs. lakhs, did {n} pay to Varghese Holdings Private Limited in FY 2024-25?",
                rp[("Varghese Holdings Private Limited", "Rent for warehouse")]),
            text("q4", "multi_hop", f"{n} pays cold storage charges to an LLP in which one of its directors is a designated partner. Which director?",
                 "Maria Varghese"),
            num("q5", "multi_hop", f"How much, in Rs. lakhs, did {n} pay in FY 2024-25 to the LLP in which its director Maria Varghese is a designated partner?",
                42.18),
            num("q6", "lookup", f"What was {n}'s other income in FY 2024-25, in Rs. lakhs?", p0["other_income"]),
        ]
    raise KeyError(s)


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DATA_DIR
    out.mkdir(parents=True, exist_ok=True)
    truth = {"description": "Synthetic companies with planted issues. All names and figures are invented.",
             "units": "Statement figures in Rs. lakhs; ledger amounts in rupees.",
             "issue_types": ISSUE_TYPES, "companies": []}
    questions: list[dict] = []
    for c in COMPANIES:
        st = compute_statements(c)
        d = out / c["slug"]
        d.mkdir(parents=True, exist_ok=True)
        md = build_markdown(c, st)
        (d / "annual_report.md").write_text(md, encoding="utf-8")
        write_pdf(c, md, d / "annual_report.pdf")
        write_xlsx(c, st, d / "financials.xlsx")
        rows, planted = build_ledger(c)
        write_ledger_csv(rows, d / "general_ledger.csv")
        for t in c["issues"]:
            assert t in ISSUE_TYPES, t
        truth["companies"].append({
            "slug": c["slug"], "name": c["name"],
            "files": sorted(p.name for p in d.iterdir()),
            "planted_issues": c["issues"],
            "decoys": c.get("decoys", []),
            "approval_limit_rupees": c["approval_limit"],
            "caro_statement_present": c["caro"],
            "ledger": {"rows": len(rows), **planted},
            "statements": {
                "pnl": {"FY 2024-25": st["pnl"][0], "FY 2023-24": st["pnl"][1]},
                "balance_sheet": {BS_CUR: st["bs"][0], BS_PRI: st["bs"][1]},
            },
            "director_interests": [{"director": a, "entity": b, "interest": x} for a, b, x in c["mbp1"]],
            "related_party_transactions": [{"party": p, "nature": nat, CUR: x, PRI: y}
                                           for p, nat, x, y in c["rp_txns"]],
        })
        questions.extend(build_questions(c, st))
        print(f"Wrote {d.relative_to(out.parent)} ({len(rows)} ledger rows, {len(md)} chars of report)")
    (out / "ground_truth.json").write_text(json.dumps(truth, indent=2) + "\n", encoding="utf-8")
    (out / "qa.json").write_text(json.dumps(questions, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(questions)} questions and the ground truth to {out}")


if __name__ == "__main__":
    main()
