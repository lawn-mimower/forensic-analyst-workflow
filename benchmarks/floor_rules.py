"""
Conventional floor for the issue-detection benchmark: a handful of simple,
deterministic rules of the kind an auditor would script in a spreadsheet.
No statistics beyond counting and no LLM.

The thresholds were fixed before any system was run. They were written by the
same author who planted the issues, so they target the planted mechanisms;
read their scores as what a targeted script achieves on this data, not as a
neutral baseline.

Rules (one per issue type):
    duplicate_entries           >= 1 group of ledger rows with the same date, account, party and amount
    round_number_entries        >= 15% of ledger amounts are multiples of Rs. 1,000
    benford_nonconformity       >= 8% of ledger amounts fall in [80%, 100%) of the approval limit
                                stated in the report (a threshold-avoidance check)
    unusual_large_payment       a single ledger amount >= 2% of revenue from operations
    ratio_anomaly               gross margin, receivables/revenue or inventories/revenue moves
                                by >= 10 percentage points year on year
    rpt_threshold_no_approval   a related-party transaction >= 10% of net worth
                                (the rule cannot read whether approval was obtained)
    rp_omitted_from_schedule    an entity from the directors' interest table (Form MBP-1) is paid in
                                the ledger but is absent from the related-party schedule
    caro_missing                revenue >= Rs. 1,000 lakhs and the text never mentions the
                                Companies (Auditor's Report) Order
    cash_loan_269ss             a sentence mentions a loan or deposit together with "in cash"
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from pathlib import Path


def _norm(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", name.lower().replace("&", " and ")).strip()


def load_ledger(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as f:
        return [{**r, "Amount": float(r["Amount"])} for r in csv.DictReader(f)]


def load_financials(path: Path) -> dict:
    """Read the three sheets of financials.xlsx into plain dicts (Rs. lakhs)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)

    def two_col(ws):
        out = {}
        for row in ws.iter_rows(values_only=True):
            if row and isinstance(row[0], str) and len(row) >= 3 and isinstance(row[1], (int, float)):
                out[row[0].strip()] = (float(row[1]), float(row[2]))
        return out

    rp = []
    for row in wb["Related Party Transactions"].iter_rows(values_only=True):
        if row and len(row) >= 5 and isinstance(row[3], (int, float)):
            rp.append({"party": row[0], "relationship": row[1], "nature": row[2],
                       "current": float(row[3]), "prior": float(row[4])})
    return {"bs": two_col(wb["Balance Sheet"]), "pnl": two_col(wb["Profit and Loss"]), "rp": rp}


def director_interest_entities(report_md: str) -> list[str]:
    """Entities from the Form MBP-1 table in the directors' report."""
    section = report_md.split("Form MBP-1", 1)[-1] if "Form MBP-1" in report_md else ""
    entities = []
    for line in section.splitlines():
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if len(cells) >= 2 and cells[0] not in ("Director", "---") and not set(cells[0]) <= set("-: "):
                entities.append(cells[1])
        elif entities and line.startswith("#"):
            break
    return entities


def approval_limit(report_md: str) -> float | None:
    m = re.search(r"payment above Rs\.?\s*([\d,]+)", report_md)
    return float(m.group(1).replace(",", "")) if m else None


def evaluate(company_dir: str | Path) -> dict:
    """Apply every rule to one company's files. Returns {flags: [...], evidence: {...}}."""
    d = Path(company_dir)
    ledger = load_ledger(d / "general_ledger.csv")
    fin = load_financials(d / "financials.xlsx")
    report = (d / "annual_report.md").read_text(encoding="utf-8")
    amounts = [r["Amount"] for r in ledger]
    n = len(amounts)
    pnl, bs = fin["pnl"], fin["bs"]
    revenue, revenue_prior = pnl["Revenue from operations"]
    ev: dict = {}
    flags: set[str] = set()

    groups = Counter((r["Date"], r["Account"], r["Party"], r["Amount"]) for r in ledger)
    dup_groups = [k for k, c in groups.items() if c > 1]
    ev["duplicate_groups"] = len(dup_groups)
    if dup_groups:
        flags.add("duplicate_entries")

    round_share = sum(1 for a in amounts if a >= 1000 and a % 1000 == 0) / n
    ev["round_thousand_share"] = round(round_share, 3)
    if round_share >= 0.15:
        flags.add("round_number_entries")

    limit = approval_limit(report)
    if limit:
        band = sum(1 for a in amounts if 0.8 * limit <= a < limit) / n
        ev["just_below_limit_share"] = round(band, 3)
        if band >= 0.08:
            flags.add("benford_nonconformity")

    largest = max(amounts)
    ev["largest_payment_pct_of_revenue"] = round(largest / (revenue * 1e5) * 100, 1)
    if largest >= 0.02 * revenue * 1e5:
        flags.add("unusual_large_payment")

    materials, materials_prior = pnl["Cost of materials consumed"]
    recv, recv_prior = bs["Trade receivables"]
    inv, inv_prior = bs["Inventories"]
    moves = {
        "gross_margin": (revenue - materials) / revenue - (revenue_prior - materials_prior) / revenue_prior,
        "receivables_to_revenue": recv / revenue - recv_prior / revenue_prior,
        "inventories_to_revenue": inv / revenue - inv_prior / revenue_prior,
    }
    ev["ratio_moves_pp"] = {k: round(v * 100, 1) for k, v in moves.items()}
    if any(abs(v) >= 0.10 for v in moves.values()):
        flags.add("ratio_anomaly")

    net_worth = bs["Equity share capital"][0] + bs["Other equity"][0]
    big = [t for t in fin["rp"] if t["current"] >= 0.10 * net_worth]
    ev["rpt_above_10pct_net_worth"] = [f"{t['party']}: {t['nature']} {t['current']:.2f}" for t in big]
    if big:
        flags.add("rpt_threshold_no_approval")

    rp_parties = {_norm(t["party"]) for t in fin["rp"]}
    ledger_parties = {_norm(r["Party"]) for r in ledger}
    missing = [e for e in director_interest_entities(report)
               if _norm(e) in ledger_parties and _norm(e) not in rp_parties]
    ev["director_entities_paid_but_not_in_rp_schedule"] = missing
    if missing:
        flags.add("rp_omitted_from_schedule")

    has_caro = "auditor's report) order" in report.lower()
    ev["caro_mentioned"] = has_caro
    if revenue >= 1000 and not has_caro:
        flags.add("caro_missing")

    cash_loan = [s.strip() for s in re.split(r"(?<=[.;])(?<!Rs\.)(?<!No\.)\s+", report)
                 if re.search(r"\b(loan|deposit)s?\b", s, re.I) and re.search(r"\bin cash\b", s, re.I)]
    ev["cash_loan_sentences"] = cash_loan
    if cash_loan:
        flags.add("cash_loan_269ss")

    return {"flags": sorted(flags), "evidence": ev}
